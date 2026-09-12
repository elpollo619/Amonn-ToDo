#!/usr/bin/env python3
# Saca de la Liste Mietvertrag neu.xlsx (hoja "Liste aktuell") los campos
# que le interesan al asistente, a JSON. Solo lectura.
import json, sys
from openpyxl import load_workbook

RUTA = "/Users/cristianamaya/Library/CloudStorage/GoogleDrive-elpollotue@gmail.com/Meine Ablage/09 Software/workpulse todo/Hans Amonn AG/Immobilien/01 Mietverträge/Liste Mietvertrag neu.xlsx"
CAMPOS = ["ObjGrp", "ObjCode", "M1VName", "M1Name", "M1Tel", "M1email",
          "Mbeginn", "Frist", "Mnetto", "Total", "Depot", "Objekt",
          "ObjektZus", "ObjAdr", "ObjOrt", "Bemerkungen"]

wb = load_workbook(RUTA, read_only=True, data_only=True)
ws = wb["Liste aktuell"]
filas = ws.iter_rows(values_only=True)
cab = [str(c).strip() if c is not None else "" for c in next(filas)]
idx = {c: cab.index(c) for c in CAMPOS if c in cab}
faltan = [c for c in CAMPOS if c not in cab]

out = []
for fila in filas:
    d = {}
    for campo, i in idx.items():
        v = fila[i] if i < len(fila) else None
        if hasattr(v, "isoformat"):
            v = v.isoformat()[:10]
        d[campo.lower()] = v
    # Una fila de verdad tiene código de objeto y algún dato del inquilino.
    if d.get("objcode") and (d.get("m1name") or d.get("total")):
        out.append(d)

print(json.dumps({"faltan": faltan, "filas": len(out), "datos": out}, ensure_ascii=False))
