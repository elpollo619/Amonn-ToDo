#!/usr/bin/env python3
# Saca los contactos de obra de las DOS listas del Drive, a JSON. Solo lectura.
#
#   1) Adressliste.xlsx      → una hoja por proyecto (G60 Muri, I16 Gampelen,
#                              770 Bremgarten, SWE A4, B4). Columnas:
#                              BKP · Arbeitsgattung · Anrede · Name-SB · Strasse ·
#                              PLZ Ort · e-mail · Telefon · Mobile
#   2) Kontaktliste_770lds_Seewer.xlsx → hoja "Kontakte" (proyecto 770). Columnas:
#                              Kategorie · BKP · Gewerk/Rolle · Firma ·
#                              Ansprechpartner · Telefon · E-Mail · Adresse/Ort ·
#                              Status · Bemerkung
#
# Los dos 770 (Adressliste y Kontaktliste) se unifican bajo un mismo proyecto.
import json
from openpyxl import load_workbook

BASE = "/Users/cristianamaya/Library/CloudStorage/GoogleDrive-elpollotue@gmail.com/Meine Ablage"
ADRESS = f"{BASE}/Adressliste.xlsx"
KONTAKT = f"{BASE}/Kontaktliste_770lds_Seewer.xlsx"
P770 = "770 Bremgarten (Seewer)"

ANREDE = {"herr", "frau", "ehegatten", "herr.", "frau.", "hr", "fr", "familie"}


def clean(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() == "none" else s


def addr(strasse, plz):
    p = [x for x in (clean(strasse), clean(plz)) if x]
    return ", ".join(p) or None


def emit(out, seen, c):
    # Un contacto necesita al menos nombre o empresa.
    if not c.get("name") and not c.get("company"):
        return
    if not c.get("name"):
        c["name"] = c["company"]
    # Dedupe por proyecto + nombre (case-insensitive).
    key = (c.get("project") or "").lower() + "|" + c["name"].lower()
    if key in seen:
        prev = seen[key]
        for k, v in c.items():          # rellena huecos del que ya estaba
            if v and not prev.get(k):
                prev[k] = v
        return
    seen[key] = c
    out.append(c)


def parse_adressliste(out, seen):
    wb = load_workbook(ADRESS, read_only=True, data_only=True)
    for sheet in wb.sheetnames:
        if sheet.lower() == "quelle":
            continue
        project = P770 if sheet.startswith("770") else sheet
        ws = wb[sheet]
        prev = None
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                continue  # cabecera
            v = [clean(x) for x in row] + [""] * 9
            bkp, role, c2, name3, strasse, plz, email, tel, mobile = v[:9]
            is_anrede = c2.lower() in ANREDE
            company = None if is_anrede else (c2 or None)
            name = name3 or ""
            # Fila de continuación (sin BKP ni Arbeitsgattung): hereda del anterior.
            if not bkp and not role and prev:
                role = prev.get("role") or ""
                if not company and not name3:
                    company = prev.get("company")
            if not name and not company:
                if role and not name3:      # solo un oficio suelto → no es contacto
                    continue
                continue
            c = {
                "name": name or (company or ""),
                "company": company,
                "role": role or None,
                "bkp": bkp or None,
                "project": project,
                "phone": tel or None,
                "mobile": mobile or None,
                "email": email or None,
                "address": addr(strasse, plz),
                "status": None,
                "notes": None,
            }
            emit(out, seen, c)
            prev = c


def parse_kontaktliste(out, seen):
    wb = load_workbook(KONTAKT, read_only=True, data_only=True)
    ws = wb["Kontakte"]
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        v = [clean(x) for x in row] + [""] * 10
        if not header_seen:
            if v[0] == "Kategorie":
                header_seen = True
            continue
        kat, bkp, role, company, name, tel, email, address, status, notes = v[:10]
        if not name and not company:
            continue
        c = {
            "name": name or company,
            "company": company or None,
            "role": role or None,
            "bkp": bkp or None,
            "project": P770,
            "phone": tel or None,
            "mobile": None,
            "email": email or None,
            "address": address or None,
            "status": status or None,
            "notes": notes or None,
        }
        emit(out, seen, c)


out, seen = [], {}
# La Kontaktliste (770) es la fuente rica: va primero para que gane en el dedupe.
parse_kontaktliste(out, seen)
parse_adressliste(out, seen)

# Casi-duplicados de las dos listas: una fila "solo empresa" (name == company)
# que duplica a una persona real de la misma empresa/proyecto. Se descarta la
# fila-placeholder y se queda la de la persona.
por_pc = {}
for c in out:
    if c.get("company"):
        por_pc.setdefault((c["project"], c["company"].lower()), []).append(c)
limpio = []
for c in out:
    comp = (c.get("company") or "").lower()
    placeholder = comp and c["name"].lower() == comp
    hay_persona = any(o is not c and o["name"].lower() != comp for o in por_pc.get((c["project"], comp), []))
    if placeholder and hay_persona:
        continue
    limpio.append(c)
out = limpio

por_proyecto = {}
for c in out:
    por_proyecto[c["project"]] = por_proyecto.get(c["project"], 0) + 1

print(json.dumps({"total": len(out), "por_proyecto": por_proyecto, "datos": out}, ensure_ascii=False))
